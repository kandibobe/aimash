"""ingest: чтение ссылок/файлов → текст для задачи агенту. Офлайн (без сети — SSRF-гард блокирует
локальные адреса до коннекта; файлы строим в памяти). Проверяем извлечение URL, SSRF-гард, HTML→текст,
парс файлов (txt/csv/json/.docx/.xlsx + отказы), и проброс контента в handle_command как ДАННЫХ.
"""

from __future__ import annotations

import io
import sys
from contextlib import contextmanager
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core import ingest  # noqa: E402


@contextmanager
def patched(obj, name, value):
    orig = getattr(obj, name)
    setattr(obj, name, value)
    try:
        yield
    finally:
        setattr(obj, name, orig)


# ── extract_urls ────────────────────────────────────────────────────────────────────
def test_extract_urls_finds_and_trims():
    text = "Глянь https://shop.ua/landing, и ещё http://example.com/x)."
    urls = ingest.extract_urls(text)
    assert urls[0] == "https://shop.ua/landing"
    assert urls[1] == "http://example.com/x"
    assert ingest.extract_urls("без ссылок") == []


# ── SSRF-гард (IP-литералы → без DNS) ────────────────────────────────────────────────
def test_is_public_host_blocks_internal():
    assert ingest._is_public_host("127.0.0.1") is False  # loopback
    assert ingest._is_public_host("10.1.2.3") is False  # private
    assert ingest._is_public_host("169.254.1.1") is False  # link-local
    assert ingest._is_public_host("0.0.0.0") is False  # unspecified
    assert ingest._is_public_host("") is False
    assert ingest._is_public_host("8.8.8.8") is True  # публичный


# ── HTML → текст ──────────────────────────────────────────────────────────────────────
def test_html_to_text_strips_and_titles():
    html = "<html><head><title>Магазин</title><style>x{}</style></head><body><h1>Купить</h1><script>bad()</script><p>быстро и дёшево</p></body></html>"
    text = ingest._html_to_text(html)
    assert "Магазин" in text and "Купить" in text and "быстро" in text
    assert "bad()" not in text and "x{}" not in text  # script/style вырезаны


# ── extract_file_text ────────────────────────────────────────────────────────────────
def test_extract_txt_csv_json():
    assert "бриф" in ingest.extract_file_text("b.txt", "бриф кампании".encode())
    assert "ключ" in ingest.extract_file_text("k.csv", "ключ,объём\nкупить,100".encode())
    assert "topic" in ingest.extract_file_text("d.json", b'{"topic": "shoes"}')


def test_extract_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["ключ", "объём"])
    ws.append(["купить телефон", 1200])
    buf = io.BytesIO()
    wb.save(buf)
    text = ingest.extract_file_text("keywords.xlsx", buf.getvalue())
    assert "купить телефон" in text and "1200" in text


def test_extract_docx():
    from docx import Document

    d = Document()
    d.add_paragraph("Бриф кампании: продаём кроссовки")
    buf = io.BytesIO()
    d.save(buf)
    text = ingest.extract_file_text("brief.docx", buf.getvalue())
    assert "Бриф кампании" in text and "кроссовки" in text


def test_extract_pdf_soft_reject():
    with pytest.raises(ingest.IngestError):
        ingest.extract_file_text("doc.pdf", b"%PDF-1.4 ...")


def test_extract_unsupported_and_empty_and_toobig():
    with pytest.raises(ingest.IngestError):
        ingest.extract_file_text("a.exe", b"MZ...")
    with pytest.raises(ingest.IngestError):
        ingest.extract_file_text("empty.txt", b"   ")
    with pytest.raises(ingest.IngestError):
        ingest.extract_file_text("big.txt", b"x" * (ingest.MAX_FILE_BYTES + 1))


def test_extract_caps_text_length():
    big = ("a" * (ingest.MAX_TEXT_CHARS + 500)).encode()
    out = ingest.extract_file_text("big.txt", big)
    assert len(out) == ingest.MAX_TEXT_CHARS


# ── fetch_url_text: гарды (офлайн) ────────────────────────────────────────────────────
async def test_fetch_rejects_non_http_scheme():
    with pytest.raises(ingest.IngestError):
        await ingest.fetch_url_text("ftp://example.com/x")


async def test_fetch_blocks_loopback_without_network():
    # event-hook _guard блокирует 127.0.0.1 ДО коннекта → IngestError, без реального запроса.
    with pytest.raises(ingest.IngestError):
        await ingest.fetch_url_text("http://127.0.0.1:9/secret")
